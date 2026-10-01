"""Document generation tools for Goal2Done."""
import os
from pathlib import Path
import re

OUTPUT_DIR = Path(os.getenv("GOAL2DONE_OUTPUT_DIR", Path.cwd() / "generated")).resolve()
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_name(name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return name or "goal2done_document"


def generate_document(format: str, title: str, content: str, filename: str = ""):
    fmt = (format or "docx").lower().lstrip(".")
    base = _safe_name(filename or title)
    if base.lower().endswith((".pdf", ".docx", ".txt", ".md")):
        base = base.rsplit(".", 1)[0]
    path = OUTPUT_DIR / f"{base}.{fmt}"

    try:
        if fmt == "txt":
            path.write_text(f"{title}\n\n{content}\n", encoding="utf-8")
        elif fmt == "md":
            path.write_text(f"# {title}\n\n{content}\n", encoding="utf-8")
        elif fmt == "docx":
            from docx import Document
            doc = Document()
            doc.add_heading(title, level=1)
            for block in content.split("\n\n"):
                doc.add_paragraph(block)
            doc.save(path)
        elif fmt == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            from xml.sax.saxutils import escape
            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(str(path), pagesize=A4)
            story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 16)]
            for block in content.split("\n\n"):
                story.append(Paragraph(escape(block).replace("\n", "<br/>"), styles["BodyText"]))
                story.append(Spacer(1, 10))
            doc.build(story)
        else:
            return {"status": "error", "type": "document", "message": "Supported formats: pdf, docx, txt, md"}

        return {
            "status": "success",
            "type": "document",
            "format": fmt,
            "title": title,
            "path": str(path),
            "filename": path.name,
        }
    except Exception as exc:
        return {"status": "error", "type": "document", "message": str(exc)}
